"""Builds GLOSSARY.md, TECHNIQUES.md and video-techniques.xlsx from the data below.
Run: pip install openpyxl && python docs/build_sheets.py
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------- glossary
GLOSSARY = ("Glossary", ["Group", "Term", "Say it like this", "Example in our films"], [
    ["Frame", "Frame", "One still image. 30 per second.", "Every frame is a pure function of time."],
    ["Frame", "FPS", "Frames per second: how smooth motion feels.", "Film 30 fps; line-boil effect at 12 fps."],
    ["Frame", "Resolution", "Pixels wide x tall.", "1080p landscape, 1080x1920 for Shorts."],
    ["Frame", "Aspect ratio", "Width : height shape.", "16:9 film, 9:16 vertical Short."],
    ["Frame", "Safe area", "Zone where text is never cut off by the phone UI.", "Lyrics kept inside it."],
    ["Composition", "Framing", "What is inside the rectangle and where.", "Tiny figure in a wide = weak framing."],
    ["Composition", "Shot types", "Wide / medium / close-up / extreme close-up.", "Close-up for the emotional line."],
    ["Composition", "Rule of thirds", "Put the subject on a third line, not the centre.", ""],
    ["Composition", "Negative space", "Empty area that gives the subject room.", "Lyric text sits there."],
    ["Composition", "Scrim", "Soft dark/light gradient behind text for contrast.", "Lyric readability on busy plates."],
    ["Composition", "Hook", "The first 1.5 s that stops the scroll.", "Every plan needs one."],
    ["Camera", "Pan / tilt", "Camera turns left-right / up-down.", ""],
    ["Camera", "Dolly", "Camera physically moves toward/away.", "2.5D dolly on still plates."],
    ["Camera", "Crane", "Camera moves up/down through space.", ""],
    ["Camera", "Rack focus", "Focus shifts from near to far subject.", "Close-up lens pass in Toota Naksha."],
    ["Camera", "Depth of field (DoF)", "Blur outside the focus plane.", "GPU bokeh particles."],
    ["Camera", "Parallax", "Near layers move more than far layers.", "Fake depth from flat plates."],
    ["Camera", "2.5D", "Flat images placed in 3D space and moved by a 3D camera.", "Toota Naksha floor-plane projection."],
    ["Layers", "Layer", "A separate stack level: background, mid, character, text.", ""],
    ["Layers", "Plate", "A background image/art layer of the scene.", "Calibrated horizon + vanishing points."],
    ["Layers", "Alpha / transparency", "The see-through part of an image.", "Sprites keyed to transparent PNG."],
    ["Layers", "Compositing", "Stacking layers into one final image.", "Canvas2D compositor."],
    ["Layers", "Blend mode", "How a layer mixes with the one below.", "Screen for glow, multiply for shadow."],
    ["Layers", "Mask / matte", "A shape that hides or reveals part of a layer.", "Line-wipe transition."],
    ["Animation", "Sprite", "A flat character image or one pose of it.", "Pose table, swim atlases."],
    ["Animation", "Sprite sheet / atlas", "Many poses packed into one image.", "Goldfish swim atlas."],
    ["Animation", "Cut-out animation", "Flat parts moved/rotated like paper puppets.", "Basement Goldfish."],
    ["Animation", "Keyframe", "A saved value at a time (position at 2.0 s).", ""],
    ["Animation", "Tweening / interpolation", "Computer fills frames between keyframes.", ""],
    ["Animation", "Easing", "Acceleration curve of a move (ease-in/out).", "Settle on the beat."],
    ["Animation", "Forward kinematics (FK)", "Pose a body by setting joint angles root-to-tip.", "Mannequin in Jod Ke Dekh."],
    ["Animation", "Rig", "Skeleton + controls that pose a model.", "Blender hand."],
    ["Animation", "Line boil", "Outlines wobble frame to frame, hand-drawn feel.", "12 fps wobble in p5."],
    ["Animation", "Crossfade ghost", "Blending two body poses makes a double image.", "Never crossfade bodies."],
    ["Animation", "Motion blur", "Smear of moving objects, like a real shutter.", "Subframe renders averaged."],
    ["3D / CGI", "CGI", "Anything made/rendered by computer graphics.", "three.js scenes, Blender props."],
    ["3D / CGI", "Mesh", "A 3D surface made of triangles.", "GLB hand."],
    ["3D / CGI", "Texture", "Image wrapped onto a mesh.", "Lyrics drawn onto a TextPlane."],
    ["3D / CGI", "Shader", "Small GPU program that decides pixel/vertex looks.", "Paper, halftone, finish shader."],
    ["3D / CGI", "Render pass", "One stage of the image pipeline.", "ScenePass, Bloom, Output."],
    ["3D / CGI", "Bloom / halation", "Bright areas glow into surroundings.", "Post finish."],
    ["3D / CGI", "Vignette / grain", "Darkened edges / film noise.", "Finish shader."],
    ["3D / CGI", "Chromatic aberration", "Colour fringes at the edges, lens flaw.", "Cinematic edge."],
    ["3D / CGI", "Particles", "Many tiny moving points (dust, sparks).", "GPU bokeh dust."],
    ["Code", "Deterministic", "Same time in = same image out. No randomness/state.", "Lets you render any frame alone."],
    ["Code", "Pure function of time", "frame = f(t) only.", "renderFilm(ctx, t)"],
    ["Code", "Canvas (HTML Canvas / Canvas2D)", "Browser surface you draw pixels onto with code.", "p5.js and the TN compositor."],
    ["Code", "WebGL / GPU", "Browser 3D drawing on the graphics card.", "three.js, shaders."],
    ["Code", "Motion Canvas", "TypeScript library: write 2D animation as generator code with a live timeline editor.", "Alternative to Remotion for vector/diagram work."],
    ["Code", "Remotion", "React framework that renders video frame by frame to MP4.", "Host for TN and JKD."],
    ["Code", "p5.js", "Friendly creative-coding library on Canvas.", "Basement Goldfish."],
    ["Code", "three.js", "JavaScript 3D engine on WebGL.", "Jod Ke Dekh."],
    ["Code", "Headless", "Runs without a visible window.", "Blender -b, puppeteer."],
    ["Code", "Concurrency", "Frames rendered in parallel.", "--concurrency=4"],
    ["Code", "Encode / mux", "Compress to MP4 / join audio and video.", "ffmpeg, CRF 17."],
    ["Audio", "BPM", "Beats per minute: sets the motion language.", "63 BPM breathes; 130 BPM cuts."],
    ["Audio", "Beat grid", "Evenly spaced beat times across the song.", "Drift-free grid."],
    ["Audio", "Onset", "The moment a sound starts.", "Per-frame onset."],
    ["Audio", "Energy curve", "Loudness over time; find section hits.", "Mark chorus drops."],
    ["Audio", "Forced alignment", "Matching known lyrics to exact audio times.", "Whisper / MMS aligner."],
    ["Audio", "Whisper", "Speech-to-text model that gives word times.", "faster-whisper large-v3."],
    ["Audio", "Stem", "One isolated track (vocals only).", "Align on the vocal stem."],
    ["Type", "Kinetic typography", "Words that move/animate with the audio.", "Lyric video."],
    ["Type", "Lane", "A fixed screen zone for one type of text.", "Main / echo / adlib lanes."],
    ["Type", "Lyric tag M / A2 / A1 / E", "Main, adlib-2, adlib-1, echo.", "Per-word tagging."],
    ["Post", "Colour grade", "Overall colour mood.", ""],
    ["Post", "Contact sheet", "Grid of frames to review a whole piece at a glance.", "tools/sheet.mjs"],
    ["Post", "Thumbnail", "The still that sells the click.", "Taken from a film frame."],
    ["Style", "Signature moment", "The one shot people remember.", "Hand signing the lyric sheet."],
    ["Style", "Callback", "Later scene echoes an earlier image.", "Ending echoes opening."],
])

# -------------------------------------------------------------- techniques
TECHNIQUES = ("Techniques", ["Technique", "What it is (1 line)", "Traditional version", "Strength", "Limit", "Tools", "Use in our code pipeline"], [
    ["Sprites", "Flat character images/poses swapped or moved", "Cel animation / cut-out puppets", "Cheap, charming, easy to swap characters", "No true 3D turn, pose limit, soft when upscaled", "p5.js, PIL/OpenCV keying, Remotion", "Basement Goldfish pose table; never crossfade bodies"],
    ["CGI", "Computer-generated 3D worlds, lit and rendered", "Pixar / VFX pipeline in Maya, Blender", "Real camera, light, depth; no art drift", "Slow render, needs modelling skill", "three.js, Blender, Remotion", "Jod Ke Dekh: mannequin + 3D set, Blender hand"],
    ["Layers", "Separate stack levels composited into one image", "Photoshop / After Effects layer stack, multiplane camera", "Control, reuse, parallax", "Layer seams; edges can halo", "Canvas2D, WebGL compositor", "Backgrounds, plates, characters, type, FX each own a layer"],
    ["Framing", "Choice of what is in shot and how big", "Cinematography / storyboard shot list", "Directs attention and emotion", "Bad framing cannot be fixed by effects", "Camera code, 2.5D projection", "Wide vs close per lyric; never type over a face"],
    ["Canvas", "Browser pixel surface you draw on with code (Canvas2D)", "Painting on a physical canvas", "Simple, fast, universal", "CPU-bound for heavy effects", "p5.js, Canvas2D API", "Compositor for plates and sprites"],
    ["Motion Canvas", "TS library, animation as generator code + live editor", "After Effects timeline, but typed", "Great for vector, diagrams, explainer feel", "2D focused; smaller ecosystem than Remotion", "motioncanvas.io", "Option for explainer/diagram sections or clean motion graphics"],
    ["Code-based animation", "Every frame computed from time by code", "Hand keyframing in a GUI", "Repeatable, exact beat sync, agent-friendly, version-controlled", "No hands-on art direction feel; setup cost", "Remotion, p5, three.js, ffmpeg", "The core of this skill: pure function of time"],
    ["Keyframe / tween", "Set values at times, computer interpolates", "Flash / After Effects keyframes", "Fast, familiar", "Looks mechanical without easing", "GSAP, Remotion interpolate", "Easing tied to beat grid"],
    ["Shaders", "GPU programs that style every pixel", "Optical filters, lab processing", "Unique looks: paper, halftone, glow", "Hard to debug; GL state quirks", "GLSL, three.js, WebGL2", "Paper slabs, halftone, finish pass"],
    ["Procedural generation", "Invent visuals from rules/noise, no art asset", "Generative art, particle sims", "Infinite variety; no source art needed", "Can feel random without direction", "Noise, particles, shaders", "Invent when you cannot generate new art"],
    ["2.5D / parallax", "Flat plates in 3D space with a moving camera", "Multiplane camera, Ken Burns", "Depth from stills", "Edges stretch past small angles", "Homography, floor-plane depth", "Toota Naksha camera moves"],
    ["Kinetic typography", "Text animated to the audio", "Title design, lyric cards", "Reads as story if lyrics live in the world", "Phone legibility traps", "Canvas textures, TextPlane", "Lyrics inside phone chat, whiteboard, invoice"],
    ["Motion graphics", "Designed shapes/graphics animated", "After Effects work", "Clean, modern, explainer", "Can look generic", "Motion Canvas, Remotion, Lottie", "Wow beats at section hits"],
    ["Rotoscoping / tracking", "Follow real footage with art or masks", "Frame-by-frame trace", "Ties art to live footage", "Labour heavy", "OpenCV, After Effects", "Use only for real-footage cutouts"],
    ["Stop-motion look", "Low fps stepped motion, handmade feel", "Clay / paper stop-motion", "Texture and charm", "Slow in real life", "Step interpolation, 12 fps boil", "Line boil, stepped limb swap"],
    ["Live action", "Real filmed footage", "Camera shoot", "Real human emotion", "Needs shoot and crew", "Camera + Python/Remotion cut", "See talking-head-director skill"],
    ["AI-generated assets", "Images/video/music from models", "Illustration, stock, composer", "Fast art and music", "Identity drift, credit honestly", "Suno, image models", "Assets are actors; declare AI use"],
])

# ---------------------------------------------------- traditional vs code pipeline
PIPELINE = ("Pipeline", ["Stage", "Traditional way", "Code-first way", "What to keep from tradition"], [
    ["1 Idea / script", "Treatment, lyric sheet, mood board", "Song facts JSON + plan doc", "Strong concept, one signature moment"],
    ["2 Storyboard", "Hand-drawn thumbnails, animatic", "Scene table with ranges, camera, light, metaphor", "Shot list thinking"],
    ["3 Design", "Concept art, character sheets", "Generated art + calibrated plates, hard rules", "Consistent style bible"],
    ["4 Layout / camera", "Layout artist sets camera", "Camera functions, 2.5D / 3D camera", "Framing and shot variety"],
    ["5 Animation", "Keyframes by hand in AE / Maya", "Pure function of time, beat-grid easing", "Timing, easing, anticipation"],
    ["6 Lighting / FX", "Compositor adds glow, particles", "Shaders, bloom, particles in the engine", "Restraint: stillness lets peaks land"],
    ["7 Audio sync", "Editor cuts to the beat by eye", "Measured BPM, onsets, aligned words", "Ear-check anyway"],
    ["8 Render", "Render farm", "Parallel frame render + ffmpeg mux", "Check the final file"],
    ["9 Review", "Dailies / director notes", "Contact sheets, strips, crops at timestamps", "Look before you say fixed"],
    ["10 Publish", "Delivery + poster", "Thumbnail from a film frame, upload pack", "Credit honestly"],
])

# -------------------------------------------------------------- recipes
RECIPES = ("Style recipes", ["Style name", "Traditional ingredient", "Code ingredient", "Result / when to use", "Done in"], [
    ["Paper puppet comedy", "Cut-out sprites", "p5 brush + line boil + beat grid", "Handmade, funny, indie tracks", "Basement Goldfish"],
    ["Cinematic ballad", "Painted plates + dolly + grading", "2.5D camera + WebGL post + audio-reactive finish", "Slow emotional songs", "Toota Naksha"],
    ["Lyrics inside the world", "Set design + props", "three.js scene table + TextPlane + Blender prop", "Rap/story songs where words are the plot", "Jod Ke Dekh"],
    ["Clean explainer lyric video", "Motion graphics", "Motion Canvas or Remotion + kinetic type", "Fast, readable, brand-safe", "Next experiment"],
    ["Sprite x CGI hybrid", "Cut-out character in a lit 3D set", "Sprite planes in three.js + shadows", "Flat character, deep world: distinct look", "Idea to try"],
    ["Procedural dream", "Abstract film / VJ visuals", "Shaders + noise + energy curve", "Instrumental or when no art exists", "Idea to try"],
    ["Live + code overlay", "Talking head shoot", "Remotion graphics + captions", "Reels, explainers", "talking-head-director"],
])

SHEETS = [GLOSSARY, TECHNIQUES, PIPELINE, RECIPES]


def md_table(title, head, rows):
    out = [f"## {title}", "", "| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    out += ["| " + " | ".join(c.replace("|", "/") for c in r) + " |" for r in rows]
    return "\n".join(out) + "\n"


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for title, head, rows in SHEETS:
        ws = wb.create_sheet(title)
        ws.append(head)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 34
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    wb.save(os.path.join(HERE, "video-techniques.xlsx"))
    open(os.path.join(HERE, "GLOSSARY.md"), "w", encoding="utf8").write(
        "# Technical words cheat sheet\n\nSay these without fumbling.\n\n" + md_table(*GLOSSARY))
    open(os.path.join(HERE, "TECHNIQUES.md"), "w", encoding="utf8").write(
        "# Techniques: traditional vs code, and how to combine them\n\n"
        + "\n".join(md_table(*s) for s in (TECHNIQUES, PIPELINE, RECIPES)))
    print("built", [s[0] for s in SHEETS])


if __name__ == "__main__":
    main()
